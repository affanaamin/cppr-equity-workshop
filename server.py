import os
import json
import datetime
from typing import Dict, Any, List, Optional
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
FILES_DIR = os.path.join(BASE_DIR, "files")
SUBMISSIONS_DIR = os.path.join(BASE_DIR, "submissions")
DATA_JSON_PATH = os.path.join(STATIC_DIR, "data.json")

os.makedirs(SUBMISSIONS_DIR, exist_ok=True)

app = FastAPI(title="CPPR Equity Budgeting Workshop Interactive Portal")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Load base data
with open(DATA_JSON_PATH, "r", encoding="utf-8") as f:
    WORKSHOP_DATA = json.load(f)

# Mount files
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/files", StaticFiles(directory=FILES_DIR), name="files")

@app.get("/", response_class=HTMLResponse)
async def read_index():
    index_file = os.path.join(STATIC_DIR, "index.html")
    with open(index_file, "r", encoding="utf-8") as f:
        return f.read()

@app.get("/api/data")
async def get_data():
    return WORKSHOP_DATA

@app.get("/api/submissions")
async def get_all_submissions():
    results = {}
    for sec in ["health", "local", "education", "finance"]:
        sec_file = os.path.join(SUBMISSIONS_DIR, f"{sec}_submission.json")
        if os.path.exists(sec_file):
            with open(sec_file, "r", encoding="utf-8") as f:
                results[sec] = json.load(f)
        else:
            results[sec] = {
                "sector": sec,
                "status": "not_started",
                "table_number": "",
                "last_saved": None,
                "answers": {}
            }
    return results

@app.get("/api/submissions/{sector}")
async def get_sector_submission(sector: str):
    sec = sector.lower().strip()
    if sec not in WORKSHOP_DATA["sectors"]:
        raise HTTPException(status_code=404, detail="Sector not found")
    sec_file = os.path.join(SUBMISSIONS_DIR, f"{sec}_submission.json")
    if os.path.exists(sec_file):
        with open(sec_file, "r", encoding="utf-8") as f:
            return json.load(f)
    return {
        "sector": sec,
        "status": "not_started",
        "table_number": "",
        "last_saved": None,
        "answers": {}
    }

@app.post("/api/submissions/{sector}")
async def save_sector_submission(sector: str, payload: Dict[str, Any]):
    sec = sector.lower().strip()
    if sec not in WORKSHOP_DATA["sectors"]:
        raise HTTPException(status_code=404, detail="Sector not found")
        
    payload["sector"] = sec
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    payload["last_saved"] = now_iso
    
    # Calculate completion
    answers = payload.get("answers", {})
    completed_count = 0
    cards_list = WORKSHOP_DATA["sectors"][sec]["cards"]
    for c in cards_list:
        cid = c["id"]
        c_ans = answers.get(cid, {})
        has_status = bool(c_ans.get("status"))
        has_evidence = bool(c_ans.get("evidence", "").strip())
        if has_status or has_evidence:
            completed_count += 1
            
    payload["completed_count"] = completed_count
    payload["total_count"] = len(cards_list)
    payload["completion_pct"] = round((completed_count / len(cards_list)) * 100) if cards_list else 0
    
    # 1. Save JSON
    json_path = os.path.join(SUBMISSIONS_DIR, f"{sec}_submission.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
        
    # 2. Auto-generate completed Excel workbook
    try:
        generate_completed_excel(sec, payload)
    except Exception as e:
        print(f"Error generating completed Excel for {sec}: {e}")
        
    # 3. Auto-generate completed Word document
    try:
        generate_completed_docx(sec, payload)
    except Exception as e:
        print(f"Error generating completed DOCX for {sec}: {e}")
        
    return {
        "status": "success",
        "message": f"Responses for {sec} saved successfully.",
        "last_saved": now_iso,
        "completed_count": completed_count,
        "completion_pct": payload["completion_pct"]
    }

@app.post("/api/reset")
async def reset_all_submissions():
    for sec in ["health", "local", "education", "finance"]:
        for prefix in [f"{sec}_submission.json", f"{sec}_Completed_Workbook.xlsx", f"{sec}_Completed_Worksheet.docx"]:
            fpath = os.path.join(SUBMISSIONS_DIR, prefix)
            if os.path.exists(fpath):
                try:
                    os.remove(fpath)
                except Exception:
                    pass
    return {"status": "success", "message": "All submissions reset"}

@app.post("/api/reset/{sector}")
async def reset_sector_submission(sector: str):
    sec = sector.lower().strip()
    for prefix in [f"{sec}_submission.json", f"{sec}_Completed_Workbook.xlsx", f"{sec}_Completed_Worksheet.docx"]:
        fpath = os.path.join(SUBMISSIONS_DIR, prefix)
        if os.path.exists(fpath):
            try:
                os.remove(fpath)
            except Exception:
                pass
    return {"status": "success", "message": f"Sector {sec} reset"}

@app.get("/api/export/{sector}/{fmt}")
async def export_sector(sector: str, fmt: str):
    sec = sector.lower().strip()
    if sec not in WORKSHOP_DATA["sectors"]:
        raise HTTPException(status_code=404, detail="Sector not found")
    
    if fmt == "excel" or fmt == "xlsx":
        fpath = os.path.join(SUBMISSIONS_DIR, f"{sec}_Completed_Workbook.xlsx")
        if not os.path.exists(fpath):
            raise HTTPException(status_code=404, detail="No submission Excel generated yet")
        return FileResponse(fpath, filename=f"{sec.capitalize()}_Completed_Workbook.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
    elif fmt == "docx" or fmt == "word":
        fpath = os.path.join(SUBMISSIONS_DIR, f"{sec}_Completed_Worksheet.docx")
        if not os.path.exists(fpath):
            raise HTTPException(status_code=404, detail="No submission Word document generated yet")
        return FileResponse(fpath, filename=f"{sec.capitalize()}_Completed_Worksheet.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        
    elif fmt == "json":
        fpath = os.path.join(SUBMISSIONS_DIR, f"{sec}_submission.json")
        if not os.path.exists(fpath):
            raise HTTPException(status_code=404, detail="No submission JSON generated yet")
        return FileResponse(fpath, filename=f"{sec}_submission.json", media_type="application/json")
        
    else:
        raise HTTPException(status_code=400, detail="Invalid format. Use xlsx, docx, or json.")

# Excel Generator Helper
def generate_completed_excel(sec: str, payload: Dict[str, Any]):
    sec_info = WORKSHOP_DATA["sectors"][sec]
    wb = openpyxl.Workbook()
    
    # Sheet 1: Overview
    ws_ov = wb.active
    ws_ov.title = "Overview & Instructions"
    ws_ov.views.sheetView[0].showGridLines = True
    
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    data_font = Font(name="Calibri", size=10)
    
    ws_ov.merge_cells("A1:H2")
    ws_ov["A1"] = f"CPPR EQUITY BUDGETING WORKSHOP\nCOMPLETED PARTICIPANT WORKBOOK — {sec_info['display_title'].upper()}"
    ws_ov["A1"].fill = title_fill
    ws_ov["A1"].font = title_font
    ws_ov["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws_ov["A4"] = "Assigned Team:"; ws_ov["A4"].font = bold_font
    ws_ov["B4"] = sec_info["display_title"]; ws_ov["B4"].font = data_font
    
    ws_ov["A5"] = "Departmental Scope:"; ws_ov["A5"].font = bold_font
    ws_ov["B5"] = sec_info["scope"]; ws_ov["B5"].font = data_font
    
    ws_ov["A6"] = "Table #:"; ws_ov["A6"].font = bold_font
    ws_ov["B6"] = payload.get("table_number", "Table # 1"); ws_ov["B6"].font = data_font
    
    ws_ov["A7"] = "Last Submitted / Saved:"; ws_ov["A7"].font = bold_font
    ws_ov["B7"] = payload.get("last_saved", ""); ws_ov["B7"].font = data_font
    
    ws_ov["A8"] = "Completion Status:"; ws_ov["A8"].font = bold_font
    ws_ov["B8"] = f"{payload.get('completed_count', 0)} of {payload.get('total_count', 14)} Indicators ({payload.get('completion_pct', 0)}%)"; ws_ov["B8"].font = bold_font
    
    ws_ov.column_dimensions["A"].width = 24
    ws_ov.column_dimensions["B"].width = 70
    
    # Sheet 2: Completed Worksheet
    ws_ws = wb.create_sheet("Equity Indicators Worksheet")
    ws_ws.views.sheetView[0].showGridLines = True
    
    ws_ws.merge_cells("A1:N2")
    ws_ws["A1"] = f"EQUITY INDICATORS EVALUATION WORKSHEET — {sec_info['display_title'].upper()}"
    ws_ws["A1"].fill = title_fill
    ws_ws["A1"].font = title_font
    ws_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
    headers = [
        ("Indicator ID", 12),
        ("Budget Stage", 20),
        ("Framework Indicator Question", 38),
        ("Assigned PDF Searched", 26),
        ("Page Reference (Doc / PDF)", 20),
        ("Exact Textual Evidence / Budget Snippet", 42),
        ("Gender", 10),
        ("Age", 10),
        ("Ethnicity / Merged", 14),
        ("Religion", 10),
        ("Disability", 11),
        ("Culture", 10),
        ("Overall Status", 14),
        ("Team Findings & Analytical Rationale", 40)
    ]
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    border_thin = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
    
    for c_idx, (h_title, w) in enumerate(headers, 1):
        c = ws_ws.cell(4, c_idx, h_title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_letter = get_column_letter(c_idx)
        ws_ws.column_dimensions[col_letter].width = w
    ws_ws.row_dimensions[4].height = 32
    
    answers = payload.get("answers", {})
    
    for r_idx, card in enumerate(sec_info["cards"], 5):
        ws_ws.row_dimensions[r_idx].height = 55
        cid = card["id"]
        c_ans = answers.get(cid, {})
        
        ws_ws.cell(r_idx, 1, cid).font = bold_font
        ws_ws.cell(r_idx, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws_ws.cell(r_idx, 1).border = border_thin
        
        ws_ws.cell(r_idx, 2, card["stage"]).font = data_font
        ws_ws.cell(r_idx, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws_ws.cell(r_idx, 2).border = border_thin
        
        ws_ws.cell(r_idx, 3, card["question"]).font = data_font
        ws_ws.cell(r_idx, 3).alignment = Alignment(wrap_text=True, vertical="center")
        ws_ws.cell(r_idx, 3).border = border_thin
        
        # User recorded answers
        c4 = ws_ws.cell(r_idx, 4, c_ans.get("pdf_searched", ""))
        c4.font = data_font; c4.alignment = Alignment(wrap_text=True, vertical="center"); c4.border = border_thin
        
        page_ref = c_ans.get("page_ref", "")
        if not page_ref and (c_ans.get("book_page") or c_ans.get("pdf_page")):
            page_ref = f"Doc p. {c_ans.get('book_page', '')} / PDF p. {c_ans.get('pdf_page', '')}"
        c5 = ws_ws.cell(r_idx, 5, page_ref)
        c5.font = data_font; c5.alignment = Alignment(horizontal="center", vertical="center"); c5.border = border_thin
        
        c6 = ws_ws.cell(r_idx, 6, c_ans.get("evidence", ""))
        c6.font = data_font; c6.alignment = Alignment(wrap_text=True, vertical="center"); c6.border = border_thin
        
        dims = c_ans.get("dimensions", [])
        dim_map = {
            7: "Gender",
            8: "Age",
            9: "Ethnicity",
            10: "Religion",
            11: "Disability",
            12: "Culture"
        }
        for d_col, d_name in dim_map.items():
            checked = any(d_name.lower() in d.lower() for d in dims)
            c_d = ws_ws.cell(r_idx, d_col, "YES" if checked else "")
            c_d.font = bold_font if checked else data_font
            c_d.alignment = Alignment(horizontal="center", vertical="center")
            c_d.border = border_thin
            if checked:
                c_d.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                
        status_val = c_ans.get("status", "")
        c13 = ws_ws.cell(r_idx, 13, status_val)
        c13.font = bold_font
        c13.alignment = Alignment(horizontal="center", vertical="center")
        c13.border = border_thin
        if status_val == "YES":
            c13.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif status_val == "NO":
            c13.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        elif status_val == "PARTIAL":
            c13.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            
        c14 = ws_ws.cell(r_idx, 14, c_ans.get("rationale", ""))
        c14.font = data_font; c14.alignment = Alignment(wrap_text=True, vertical="center"); c14.border = border_thin
        
    out_xlsx = os.path.join(SUBMISSIONS_DIR, f"{sec}_Completed_Workbook.xlsx")
    wb.save(out_xlsx)
    print(f"Saved completed Excel: {out_xlsx}")

# Word Generator Helper
def generate_completed_docx(sec: str, payload: Dict[str, Any]):
    sec_info = WORKSHOP_DATA["sectors"][sec]
    doc = docx.Document()
    
    # Page setup (no headers, no footers)
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs: p.text = ""
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs: p.text = ""
    
    # Title Banner
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(4)
    p_t.paragraph_format.space_after = Pt(1)
    r_t = p_t.add_run("CPPR EQUITY BUDGETING WORKSHOP")
    r_t.font.name = "Calibri"; r_t.font.size = Pt(20); r_t.bold = True
    r_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    r_s = p_sub.add_run(f"COMPLETED PARTICIPANT WORKSHEET — {sec_info['display_title'].upper()}")
    r_s.font.name = "Calibri"; r_s.font.size = Pt(11); r_s.bold = True
    r_s.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(6)
    r_m = p_meta.add_run(f"Departmental Focus: {sec_info['scope']}  |  Table: {payload.get('table_number', '1')}  |  Submitted: {payload.get('last_saved', '')}")
    r_m.font.name = "Calibri"; r_m.font.size = Pt(9); r_m.font.italic = True
    r_m.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    
    answers = payload.get("answers", {})
    
    # 14 Cards
    for idx, card in enumerate(sec_info["cards"], 1):
        cid = card["id"]
        c_ans = answers.get(cid, {})
        
        tbl = doc.add_table(rows=6, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Border XML
        tblPr = tbl._tbl.tblPr
        tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'</w:tblBorders>'
        ))
        
        # Row 0: Banner
        c0 = tbl.cell(0, 0); c0.merge(tbl.cell(0, 1))
        c0._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79"/>'))
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_before = Pt(0); p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(f"INDICATOR {cid}: {card['stage'].upper()}")
        r0.bold = True; r0.font.name = "Calibri"; r0.font.size = Pt(10); r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        # Row 1: Question
        c1 = tbl.cell(1, 0); c1.merge(tbl.cell(1, 1))
        c1._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F5F9"/>'))
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_before = Pt(0); p1.paragraph_format.space_after = Pt(1)
        r1_lbl = p1.add_run("Framework Indicator Question: "); r1_lbl.bold = True; r1_lbl.font.size = Pt(9); r1_lbl.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        p1.add_run(card["question"]).font.size = Pt(9)
        
        p1_src = c1.add_paragraph(); p1_src.paragraph_format.space_before = Pt(1); p1_src.paragraph_format.space_after = Pt(0)
        r1_src = p1_src.add_run("Prescribed Source Documents: "); r1_src.bold = True; r1_src.font.size = Pt(8.5); r1_src.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        r1_stxt = p1_src.add_run(card["prescribed_sources"]); r1_stxt.font.size = Pt(8.5); r1_stxt.italic = True
        
        # Row 2: PDF Searched & Page Reference
        c20 = tbl.cell(2, 0); c21 = tbl.cell(2, 1)
        c20.width = Inches(4.2); c21.width = Inches(3.2)
        p20 = c20.paragraphs[0]; p20.paragraph_format.space_before = Pt(0); p20.paragraph_format.space_after = Pt(0)
        p20.add_run("Assigned PDF Searched:\n").bold = True
        p20.add_run(c_ans.get("pdf_searched", "[None recorded]")).font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        
        p21 = c21.paragraphs[0]; p21.paragraph_format.space_before = Pt(0); p21.paragraph_format.space_after = Pt(0)
        p21.add_run("Page Reference (Dual Citation):\n").bold = True
        page_val = c_ans.get("page_ref", "")
        if not page_val and (c_ans.get("book_page") or c_ans.get("pdf_page")):
            page_val = f"Doc p. {c_ans.get('book_page', '')} / PDF p. {c_ans.get('pdf_page', '')}"
        p21.add_run(page_val if page_val else "[None recorded]").font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        
        # Row 3: Evidence
        c3 = tbl.cell(3, 0); c3.merge(tbl.cell(3, 1))
        p3 = c3.paragraphs[0]; p3.paragraph_format.space_before = Pt(0); p3.paragraph_format.space_after = Pt(1)
        p3.add_run("Exact Textual Evidence & Monetary Figures:\n").bold = True
        ev_text = c_ans.get("evidence", "").strip()
        p3.add_run(ev_text if ev_text else "[No textual evidence entered]")
        
        # Row 4: Dimensions & Status
        c40 = tbl.cell(4, 0); c41 = tbl.cell(4, 1)
        c40.width = Inches(4.2); c41.width = Inches(3.2)
        c40._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>'))
        c41._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>'))
        
        p40 = c40.paragraphs[0]; p40.paragraph_format.space_before = Pt(0); p40.paragraph_format.space_after = Pt(1)
        p40.add_run("Disaggregated Dimensions:\n").bold = True
        dims = c_ans.get("dimensions", [])
        if dims:
            p40.add_run(", ".join(dims)).font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        else:
            p40.add_run("[None checked]")
            
        p41 = c41.paragraphs[0]; p41.paragraph_format.space_before = Pt(0); p41.paragraph_format.space_after = Pt(1)
        p41.add_run("Overall Status Assessment:\n").bold = True
        st_val = c_ans.get("status", "[Not Assessed]")
        r_st = p41.add_run(f"[{st_val}]")
        r_st.bold = True
        if st_val == "YES": r_st.font.color.rgb = RGBColor(0x27, 0x7E, 0x48)
        elif st_val == "NO": r_st.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif st_val == "PARTIAL": r_st.font.color.rgb = RGBColor(0xED, 0x7D, 0x31)
        
        # Row 5: Rationale
        c5 = tbl.cell(5, 0); c5.merge(tbl.cell(5, 1))
        p5 = c5.paragraphs[0]; p5.paragraph_format.space_before = Pt(0); p5.paragraph_format.space_after = Pt(1)
        p5.add_run("Team Findings & Analytical Rationale:\n").bold = True
        rat_text = c_ans.get("rationale", "").strip()
        p5.add_run(rat_text if rat_text else "[No analytical rationale entered]")
        
        for r in tbl.rows:
            r._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
        if idx % 2 == 0 and idx < len(sec_info["cards"]):
            doc.add_page_break()
        else:
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            
    out_docx = os.path.join(SUBMISSIONS_DIR, f"{sec}_Completed_Worksheet.docx")
    doc.save(out_docx)
    print(f"Saved completed DOCX: {out_docx}")

if __name__ == "__main__":
    import uvicorn
    print("Starting CPPR Equity Budgeting Workshop Web Application...")
    print("Access URL: http://localhost:8000")
    print("Facilitator Dashboard: http://localhost:8000/?sector=admin")
    uvicorn.run(app, host="0.0.0.0", port=8000)
